# /// script
# requires-python = ">=3.11"
# dependencies = ["openpyxl==3.1.5"]
# ///
"""Independently compare every CSV cell with the workbook using openpyxl."""

import argparse
import csv
import gzip
import hashlib
import struct
from pathlib import Path

from openpyxl import load_workbook


def verify(source: Path, data: Path) -> None:
    if hashlib.sha256(source.read_bytes()).hexdigest() != (
        "8b532fe953400725c5e5e13541adf984652954ff7ee8c71618381eddd9ada717"
    ):
        raise ValueError("Source workbook differs from the reviewed bytes.")
    workbook = load_workbook(source, read_only=True, data_only=False)
    if workbook.sheetnames != ["Sheet1"]:
        raise ValueError("Unexpected source sheets.")
    expected = list(workbook["Sheet1"].values)
    workbook.close()
    with gzip.open(data, "rt", encoding="utf-8", newline="") as stream:
        actual = list(csv.reader(stream))
    if len(expected) != len(actual) or len(actual) != 236:
        raise ValueError("Row counts differ.")
    for i, (source_row, csv_row) in enumerate(zip(expected, actual, strict=True)):
        if len(source_row) != len(csv_row) or len(csv_row) != 140:
            raise ValueError(f"Row {i} widths differ.")
        for j, (source_value, csv_value) in enumerate(
            zip(source_row, csv_row, strict=True)
        ):
            if i == 0 or j == 1:
                matches = source_value == csv_value
            elif j == 0:
                matches = source_value == int(csv_value)
            else:
                matches = struct.pack("<d", source_value) == struct.pack(
                    "<d", float(csv_value)
                )
            if not matches:
                raise ValueError(f"Workbook and CSV differ at row {i}, column {j}.")
    print(
        "All 32,900 data cells, headers, and numeric binary64 values "
        "match the workbook."
    )


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("source", type=Path)
    parser.add_argument("data", type=Path)
    args = parser.parse_args()
    verify(args.source, args.data)
